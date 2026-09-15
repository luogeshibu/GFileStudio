from pathlib import Path
import xml.etree.ElementTree as ET

from openpyxl import load_workbook

from g_file_studio.models import InputMode
from g_file_studio.services.site_profile_service import SiteSmartProfile
from g_file_studio.services.symbol_inventory_service import (
    extract_file_inventory,
    infer_symbol_usage,
    process_symbol_inventory,
)


def _profile() -> SiteSmartProfile:
    devref = "#Transformer_OH.pb.icn.g:Transformer_OH"
    return SiteSmartProfile(
        profile_name="inventory-test",
        site_name="TEST",
        smart_lbs_devref="",
        smart_breaker_devref="",
        custom_symbols=[{
            "uid": "tr",
            "scope": "ANY",
            "role": "Transformer",
            "device_type": "Transformer",
            "device_subtype": "OH",
            "device_level": "独立设备",
            "symbol_usage": "设备",
            "element_tag": "TransformerDis",
            "standard_devref": devref,
            "match_attr": "XML元素",
            "match_value": "",
            "source_file": "Transformer_OH.pb.icn.g",
        }],
        symbol_catalog={
            devref: {
                "devref": devref,
                "element_tag": "TransformerDis",
                "element_id": "Transformer_OH",
                "source_file": "Transformer_OH.pb.icn.g",
            }
        },
    )


def _business_g(path: Path) -> None:
    root = ET.Element("G", {"facID": "1", "facName": "TEST F01"})
    layer = ET.SubElement(root, "Layer")
    ET.SubElement(layer, "Text", {"id": "t1", "x": "90", "y": "20", "w": "40", "h": "20", "ts": "97803"})
    ET.SubElement(layer, "TransformerDis", {
        "id": "d1", "x": "100", "y": "70", "w": "40", "h": "40",
        "devref": "#Transformer_OH.pb.icn.g:Transformer_OH",
    })
    ET.SubElement(layer, "Text", {"id": "t2", "x": "290", "y": "20", "w": "50", "h": "20", "ts": "972265"})
    ET.SubElement(layer, "TransformerDis", {
        "id": "d2", "x": "300", "y": "70", "w": "40", "h": "40",
        "devref": "#Transformer_OLD.pb.icn.g:Transformer_OLD",
    })
    ET.ElementTree(root).write(path, encoding="utf-8", xml_declaration=True)


def test_status_icon_suffix_is_initially_classified_as_status():
    assert infer_symbol_usage(source_file="NariPd_Generator.zt.icn.g") == "状态图元"


def test_extract_all_configured_symbols_names_and_standard_status(tmp_path: Path):
    source = tmp_path / "business.sln.pic.g"
    _business_g(source)
    symbols, devices, unmapped, warnings = extract_file_inventory(source, _profile())
    assert not warnings
    assert not unmapped
    assert len(symbols) == 2
    assert len(devices) == 2
    assert [row["实例名称"] for row in devices] == ["97803", "972265"]
    assert devices[0]["标准校验"] == "PASS"
    assert devices[1]["标准校验"] == "MISMATCH"
    assert "标准图元引用不一致" in devices[1]["异常类型"]


def test_export_device_comparison_excel(tmp_path: Path):
    source = tmp_path / "business.sln.pic.g"
    _business_g(source)
    out = tmp_path / "out"
    result = process_symbol_inventory(source, InputMode.SINGLE_FILE, out, _profile())
    workbook = Path(result.statistics["excel_path"])
    assert workbook.exists()
    wb = load_workbook(workbook, read_only=True, data_only=True)
    required = {"解析概览", "文件汇总", "全部设备", "设备类型统计", "设备对比", "全部图元实例", "全部内容清单", "XML结构统计", "图元校验异常", "未定义图元", "图元统计"}
    assert required.issubset(set(wb.sheetnames))
    ws = wb["设备对比"]
    rows = list(ws.iter_rows(values_only=True))
    header = list(rows[0])
    name_idx = header.index("实例名称")
    status_idx = header.index("标准校验")
    assert [rows[1][name_idx], rows[2][name_idx]] == ["97803", "972265"]
    assert [rows[1][status_idx], rows[2][status_idx]] == ["PASS", "MISMATCH"]
