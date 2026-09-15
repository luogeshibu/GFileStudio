from openpyxl import load_workbook

from g_file_studio.services.symbol_inventory_service import (
    _device_form,
    _is_primary_device_row,
    _migration_device_row,
    _style_workbook,
)


def _inside_rmu_independent_row():
    return {
        "GFile": "AJWD-30.g", "facID": "", "facName": "",
        "图元用途": "设备", "设备类型": "NariPd_Earth", "迁移设备类型": "NARIPD_EARTH",
        "设备子类型": "", "设备层级": "独立设备",  # legacy/profile value
        "SMART/NORMAL": "SMART", "实例名称": "", "所属RMU": "32954", "RMU类型": "3L1T",
        "XML元素": "pwbh", "ElementID": "182000026", "实际devref": "#NariPd_Earth.pwbh.icn.g:NariPd_Earth",
        "标准校验": "PASS", "名称来源": "", "名称置信度": "LOW",
    }


def test_parent_rmu_overrides_legacy_standalone_device_form():
    row = _inside_rmu_independent_row()
    assert _device_form(row) == "RMU_COMPONENT"
    migration = _migration_device_row(row)
    assert migration["ParentRMU"] == "32954"
    assert migration["DeviceForm"] == "RMU_COMPONENT"
    assert not _is_primary_device_row(row)


def test_rmu_internal_rows_are_kept_in_detail_but_excluded_from_main_and_main_statistics(tmp_path):
    rows = [
        {
            "GFile": "AJWD-30.g", "facID": "", "facName": "", "图元用途": "设备", "设备类型": "RMU",
            "迁移设备类型": "RMU", "设备子类型": "3L1T", "设备层级": "组合设备", "SMART/NORMAL": "SMART",
            "实例名称": "32954", "所属RMU": "", "RMU类型": "3L1T", "XML元素": "组合识别", "ElementID": "2000046",
            "标准校验": "COMPOSITE", "名称来源": "identify_rmus()", "名称置信度": "HIGH",
        },
        {
            "GFile": "AJWD-30.g", "facID": "", "facName": "", "图元用途": "设备", "设备类型": "Transformer",
            "迁移设备类型": "TRANSFORMER", "设备子类型": "", "设备层级": "独立设备", "SMART/NORMAL": "NORMAL",
            "实例名称": "99968", "所属RMU": "", "RMU类型": "", "XML元素": "TransformerDis", "ElementID": "115000048",
            "实际devref": "#NariPd_Transformer_OH.pb.icn.g:NariPd_Transformer_OH", "标准校验": "PASS",
            "名称来源": "Nearby Text", "名称置信度": "HIGH",
        },
        _inside_rmu_independent_row(),
    ]
    out = tmp_path / "inventory.xlsx"
    _style_workbook(out, [], rows, [], [], content_rows=[], file_summaries=[])
    wb = load_workbook(out, read_only=True, data_only=True)

    # Header + 3 detail rows; ParentRMU row is retained for audit/migration details.
    assert wb["ADMS-SLD设备明细"].max_row == 4
    # Header + RMU + Transformer only.  The object inside RMU is not a main device.
    assert wb["ADMS-SLD主设备"].max_row == 3
    # Header + the contained object.
    assert wb["RMU内部设备"].max_row == 2

    stats = list(wb["迁移设备类型统计"].iter_rows(min_row=2, values_only=True))
    types = {row[0]: row[1] for row in stats}
    assert types == {"RMU": 1, "TRANSFORMER": 1}
