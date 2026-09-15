import xml.etree.ElementTree as ET

from openpyxl import load_workbook

from g_file_studio.services.symbol_inventory_service import (
    _migration_device_row,
    _migration_device_type_from_values,
    _resolve_instance_name,
    _style_workbook,
)


def test_migration_device_type_normalizes_supported_business_families():
    assert _migration_device_type_from_values(device_type="Transformer") == "TRANSFORMER"
    assert _migration_device_type_from_values(actual_devref="#Fuse_NON_SMART.zwk.icn.g:Fuse_NON_SMART") == "FUSE"
    assert _migration_device_type_from_values(actual_devref="#RMU_LBS_S.zwk.icn.g:RMU_LBS_S") == "LBS"
    assert _migration_device_type_from_values(device_type="Recloser", actual_devref="#AR_S.zwk.icn.g:AR_S") == "REC"
    # AR_* alone is intentionally not hard-coded to REC; the user-confirmed profile
    # remains authoritative for ambiguous site naming families.
    assert _migration_device_type_from_values(actual_devref="#AR_S_H.zwk.icn.g:AR_S_H") != "REC"
    assert _migration_device_type_from_values(device_type="SFI") == "SFI"
    assert _migration_device_type_from_values(actual_devref="#Circuit_Breaker_NO-SMART.zwk.icn.g:Circuit_Breaker_NO-SMART") == "CB"


def test_migration_row_exposes_name_smart_parent_and_quality():
    row = _migration_device_row({
        "GFile": "A.g",
        "设备类型": "LBS",
        "设备子类型": "",
        "实例名称": "Y1",
        "SMART/NORMAL": "SMART",
        "设备层级": "设备内部部件",
        "所属RMU": "32954",
        "RMU类型": "3L1T",
        "XML元素": "CBreakerDis",
        "ElementID": "117000001",
        "实际devref": "#RMU_LBS_S.zwk.icn.g:RMU_LBS_S",
        "标准校验": "PASS",
        "名称来源": "p_NameString",
        "名称置信度": "HIGH",
    })
    assert row["DeviceType"] == "LBS"
    assert row["DeviceName"] == "Y1"
    assert row["IsSmart"] == "YES"
    assert row["SmartType"] == "SMART"
    assert row["DeviceForm"] == "RMU_COMPONENT"
    assert row["ParentRMU"] == "32954"
    assert row["QualityStatus"] == "PASS"


def test_fuse_fc_annotation_is_not_used_as_migration_device_name():
    fuse = ET.Element("CBreakerDis", {"x": "100", "y": "100", "w": "48", "h": "32"})
    fc = ET.Element("Text", {"ts": "F.C", "x": "95", "y": "90", "w": "67", "h": "50"})
    name, source, confidence = _resolve_instance_name(
        fuse,
        device_type="Fuse",
        rmu_name="",
        texts=[fc],
        consumed_text_ids=set(),
    )
    assert name == ""
    assert confidence == "LOW"


def test_multisheet_adms_sld_workbook_contains_stable_device_sheets(tmp_path):
    rows = [
        {
            "GFile": "A.g", "facID": "", "facName": "", "图元用途": "设备", "设备类型": "RMU",
            "迁移设备类型": "RMU", "设备子类型": "3L1T", "设备层级": "组合设备", "设备形态": "COMPOSITE",
            "SMART/NORMAL": "SMART", "是否智能": "YES", "智能类型": "SMART", "智能判断来源": "RMU Context",
            "实例名称": "32954", "所属RMU": "", "RMU类型": "3L1T", "XML元素": "组合识别", "ElementID": "2000046",
            "标准校验": "COMPOSITE", "名称来源": "identify_rmus()", "名称置信度": "HIGH",
        },
        {
            "GFile": "A.g", "facID": "", "facName": "", "图元用途": "设备", "设备类型": "Transformer",
            "迁移设备类型": "TRANSFORMER", "设备子类型": "", "设备层级": "独立设备", "设备形态": "STANDALONE",
            "SMART/NORMAL": "NORMAL", "是否智能": "NO", "智能类型": "NORMAL", "智能判断来源": "Profile/Context",
            "实例名称": "99968", "所属RMU": "", "RMU类型": "", "XML元素": "TransformerDis", "ElementID": "115000048",
            "实际devref": "#NariPd_Transformer_OH.pb.icn.g:NariPd_Transformer_OH", "标准校验": "PASS",
            "名称来源": "Nearby Text", "名称置信度": "HIGH",
        },
        {
            "GFile": "A.g", "facID": "", "facName": "", "图元用途": "设备组成图元", "设备类型": "LBS",
            "迁移设备类型": "LBS", "设备子类型": "", "设备层级": "设备内部部件", "设备形态": "RMU_COMPONENT",
            "SMART/NORMAL": "SMART", "是否智能": "YES", "智能类型": "SMART", "智能判断来源": "Profile/Context",
            "实例名称": "Y1", "所属RMU": "32954", "RMU类型": "3L1T", "XML元素": "CBreakerDis", "ElementID": "117000039",
            "实际devref": "#RMU_LBS_S.zwk.icn.g:RMU_LBS_S", "标准校验": "PASS",
            "名称来源": "p_NameString", "名称置信度": "HIGH",
        },
    ]
    out = tmp_path / "g-content-inventory.xlsx"
    _style_workbook(out, [], rows, [], [], content_rows=[], file_summaries=[])
    wb = load_workbook(out, read_only=True)
    expected = {
        "ADMS-SLD设备明细", "ADMS-SLD主设备", "迁移设备类型统计",
        "RMU", "TRANSFORMER", "LBS", "FUSE", "REC", "SFI", "CB",
        "RMU内部设备", "未命名设备", "其他设备",
    }
    assert expected.issubset(set(wb.sheetnames))
    assert wb["ADMS-SLD设备明细"].max_row == 4
    assert wb["ADMS-SLD主设备"].max_row == 3
    assert wb["RMU内部设备"].max_row == 2
    assert wb["TRANSFORMER"].max_row == 2
    assert wb["LBS"].max_row == 2
