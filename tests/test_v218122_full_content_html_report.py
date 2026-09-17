from pathlib import Path
import xml.etree.ElementTree as ET

from openpyxl import load_workbook

from g_file_studio.models import InputMode
from g_file_studio.services.site_profile_service import SiteSmartProfile
from g_file_studio.services.symbol_inventory_service import process_symbol_inventory


def _profile() -> SiteSmartProfile:
    tr = "#Transformer_OH.pb.icn.g:Transformer_OH"
    sw = "#Load_Breaker_Switch_NON-SMART.zwk.icn.g:Load_Breaker_Switch_NON-SMART"
    return SiteSmartProfile(
        profile_name="full-inventory",
        site_name="Jeddah",
        profile_version=12,
        smart_lbs_devref="",
        smart_breaker_devref="",
        custom_symbols=[
            {
                "uid": "tr", "scope": "ANY", "role": "Transformer", "symbol_usage": "设备",
                "device_type": "Transformer", "device_subtype": "OH", "device_level": "独立设备",
                "element_tag": "TransformerDis", "standard_devref": tr, "match_attr": "XML元素", "match_value": "",
                "source_file": "Transformer_OH.pb.icn.g",
            },
            {
                "uid": "sw", "scope": "ANY", "role": "LBS", "symbol_usage": "设备",
                "device_type": "LBS", "device_subtype": "Outdoor", "device_level": "独立设备",
                "element_tag": "CBreakerDis", "standard_devref": sw, "match_attr": "XML元素", "match_value": "",
                "source_file": "Load_Breaker_Switch_NON-SMART.zwk.icn.g",
            },
        ],
        symbol_catalog={
            tr: {"devref": tr, "element_tag": "TransformerDis", "element_id": "Transformer_OH", "source_file": "Transformer_OH.pb.icn.g"},
            sw: {"devref": sw, "element_tag": "CBreakerDis", "element_id": "Load_Breaker_Switch_NON-SMART", "source_file": "Load_Breaker_Switch_NON-SMART.zwk.icn.g"},
        },
    )


def _g(path: Path) -> None:
    root = ET.Element("G", {"facID": "ABH", "facName": "ABH-12"})
    layer = ET.SubElement(root, "Layer", {"id": "L1"})
    ET.SubElement(layer, "Text", {"id": "t1", "x": "90", "y": "20", "w": "60", "h": "20", "ts": "TR-01"})
    ET.SubElement(layer, "TransformerDis", {"id": "tr1", "x": "100", "y": "70", "w": "40", "h": "40", "devref": "#Transformer_OH.pb.icn.g:Transformer_OH"})
    ET.SubElement(layer, "Text", {"id": "t2", "x": "280", "y": "20", "w": "60", "h": "20", "ts": "SW-01"})
    ET.SubElement(layer, "CBreakerDis", {"id": "sw1", "x": "300", "y": "70", "w": "28", "h": "30", "devref": "#Load_Breaker_Switch_NON-SMART.zwk.icn.g:Load_Breaker_Switch_NON-SMART"})
    ET.SubElement(layer, "ConnectLine", {"id": "c1", "sourceId": "tr1", "targetId": "sw1"})
    ET.SubElement(layer, "PokeDis", {"id": "p1", "ahref": "ABH-13"})
    group = ET.SubElement(layer, "Group", {"id": "g1"})
    ET.SubElement(group, "AnalogSignal", {"id": "a1", "value": "12.3"})
    ET.ElementTree(root).write(path, encoding="utf-8", xml_declaration=True)


def test_full_content_excel_and_html_are_generated(tmp_path: Path):
    source = tmp_path / "ABH-12.sln.pic.g"
    _g(source)
    result = process_symbol_inventory(source, InputMode.SINGLE_FILE, tmp_path / "out", _profile())

    excel = Path(result.statistics["excel_path"])
    html = Path(result.statistics["html_path"])
    assert excel.exists()
    assert html.exists()
    assert result.statistics["业务设备总数"] == 2
    assert result.statistics["设备类型数"] == 2
    assert result.statistics["全部XML对象"] >= 9

    wb = load_workbook(excel, read_only=True, data_only=True)
    for sheet in ("全部设备", "设备类型统计", "全部内容清单", "XML结构统计", "连接拓扑", "Poke跳转", "量测信号"):
        assert sheet in wb.sheetnames
    device_rows = list(wb["全部设备"].iter_rows(values_only=True))
    header = list(device_rows[0])
    type_idx = header.index("设备类型")
    assert {row[type_idx] for row in device_rows[1:]} == {"Transformer", "LBS"}

    html_text = html.read_text(encoding="utf-8")
    assert "G 图形内容解析报告" in html_text
    assert "全部业务设备" in html_text
    assert "Transformer" in html_text
    assert "完整 G XML / 对象清单" in html_text
    assert "ConnectLine" in html_text
    assert "AnalogSignal" in html_text


def test_global_content_analysis_page_is_removed():
    assert not Path("g_file_studio/ui/pages/symbol_inventory_page.py").exists()
